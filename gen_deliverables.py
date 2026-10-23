#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génère les livrables ARCA v2 en .docx."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "C:/Users/Administrateur/AppData/Local/hermes/hermes-agent/workspace/ARCA/DELIVERABLES_ARCA_v2.docx"

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Cm(29.7); sec.page_height = Cm(21.0)
sec.top_margin = Cm(1.5); sec.bottom_margin = Cm(1.5)
sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'; normal.font.size = Pt(10)

def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),hex_color)
    tcPr.append(shd)

def set_width(cell, cm):
    cell.width = Cm(cm)

def set_cell_margins(cell, top=30, bottom=30, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for tag,val in (('top',top),('bottom',bottom),('start',left),('end',right)):
        e = OxmlElement(f'w:{tag}'); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); m.append(e)
    tcPr.append(m)

def vcenter(cell):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr(); th = OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)

def add_para(text, size=10, bold=False, italic=False, color=None, after=4, before=0, align=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(before)
    if align: p.alignment = align
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return p

def add_table(headers, rows, widths, font_size=8.5, header_color="1F3864"):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for j,h in enumerate(headers):
        set_width(hdr[j], widths[j]); shade(hdr[j], header_color)
        p = hdr[j].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.bold = True; r.font.size = Pt(font_size); r.font.color.rgb = RGBColor.from_string("FFFFFF")
        set_cell_margins(hdr[j], top=20, bottom=20); vcenter(hdr[j])
    repeat_header(t.rows[0])
    for i,row in enumerate(rows):
        cells = t.add_row().cells
        for j,val in enumerate(row):
            set_width(cells[j], widths[j]); set_cell_margins(cells[j])
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(val)); r.font.size = Pt(font_size)
            if i%2==1: shade(cells[j], "F2F7FB")
            vcenter(cells[j])
    return t

# ============================ TITRE ============================
add_para("DELIVERABLES — ARCA v2.0", size=20, bold=True, color="1F3864", after=2)
add_para("Agent de Référentiel et de Construction · Version intégrée · 22 octobre 2026",
         size=11, color="2E5496", after=8)

# ============================ VÔTEL 1 — MANDAT ============================
add_para("VÔTEL 1/2 — Mandat", size=12, bold=True, color="1F3864", after=4)
add_table(
    ["Champ", "Contenu"],
    [
        ("Périmètre normatif", "ISO 9001 · 14001 · 45001 · 55001 · 50001 · 22301 · 37001 · 37301 · OACI Annexes 14/19, Doc 9859 · ACA niveaux 1 à 5"),
        ("Rôle", "Architecte documentaire et ingénieur de la preuve d'un SM intégré · trois casquettes : expert normatif, ingénieur documentaire, intégrateur de compétences (clause 7.2)"),
        ("Mission", "Transformer toute exigence — normative, réglementaire aéronautique, sectorielle, interne — en un corpus unique et auditable, puis en dispositif de compétence et de preuve tenant devant un organisme certificateur, une autorité de l'aviation civile et un vérificateur carbone"),
        ("Objectif (6 à 9 mois)", "60–90 documents à 100 % · matrice exigence → document → preuve sans exigence orpheline · 99 compétences + 26 modules · dossier ACA déposé · zéro écart critique en audit à blanc, zéro réserve en vérification carbone · complétude ≥ 95 %"),
        ("Contexte à cadrer", "Statut de l'aérodrome et catégorie RFFS · référentiels réellement visés · périmètre (siège, plateforme, filiales, tiers) · éditions et échéances · maturité · outils (GED, GRC, CMMS, EMS, ECCAIRS, outil carbone, LMS) · sponsor exécutif · contraintes de confidentialité et données personnelles"),
    ],
    [3.0, 18.5], font_size=8.5
)

# ============================ GUIDE DE PRODUCTION ============================
add_para("Guide de production — 7 étapes", size=12, bold=True, color="1F3864", after=4)
add_table(
    ["Étape", "Contenu", "Livrable de sortie"],
    [
        ("1. Cadrage", "Périmètre, éditions, statut de l'aérodrome, maturité, outils, sponsor, contraintes", "Note de cadrage validée (2–3 pages)"),
        ("2. Architecture", "Table de correspondance HLS, découpage documentaire, propriétaires, cycle de vie", "Architecture documentaire + liste nominique"),
        ("3. Analyse d'écarts", "Exigence par exigence, preuve existante, écart, criticité, action", "Registre des écarts priorisé"),
        ("4. Rédaction", "Documents en brouillon, citation de clause et renvoi à la preuve", "Corpus brouillon versionné"),
        ("5. Compétences", "Traduction des exigences en compétences, niveaux cibles, plan de formation", "Matrice 99 lignes + plan 26 modules"),
        ("6. Mise en preuve", "Registres, enregistrements, dossier d'audit, dossier ACA", "Dossier de preuves indexé"),
        ("7. Test", "Audit à blanc, simulation de vérification carbone, revue de direction", "Rapport d'audit à blanc + plan d'action"),
    ],
    [2.6, 9.0, 9.9], font_size=8.5
)

# ============================ GRILLE DE LIRE ============================
add_para("Grille de lecture / checklist de succès", size=12, bold=True, color="1F3864", after=4)
add_table(
    ["Indicateur", "Cible", "Statut à remplir"],
    [
        ("Couverture des informations documentées exigées (11 référentiels)", "100 %, zéro exigence orpheline", ""),
        ("Documents validés et diffusés", "≥ 95 % du corpus", ""),
        ("Matrice exigence → document → preuve à jour", "100 %, revue trimestrielle", ""),
        ("Dispositif clause 7.2 opérationnel", "99 compétences, écarts calculés, 26 modules", ""),
        ("Audit à blanc", "Zéro écart critique", ""),
        ("Vérification carbone", "Zéro réserve majeure", ""),
        ("Dossier ACA", "Déposé, niveau visé accepté", ""),
        ("Révision de direction", "Tenue avec décisions tracées", ""),
    ],
    [8.5, 6.0, 7.0], font_size=8.5
)

# ============================ MATRICE 7.2 — 13 DOMAINES ============================
add_para("Dispositif clause 7.2 — 13 domaines de compétence", size=12, bold=True, color="1F3864", after=4)
add_para("Répartition des 99 compétences. Le détail (critère observable, niveau cible, preuve, module) figure dans le classeur Excel associé.",
         size=8.5, italic=True, color="7c8899", after=4)
add_table(
    ["Domaine", "Nb", "Produit par l'agent", "Reste humain"],
    [
        ("A. Socle des systèmes de management", "12", "Architecture intégrée, table HLS, programme d'audit, revue de direction, registre des risques", "Décider des priorités et allouer les ressources"),
        ("B. Qualité — ISO 9001", "4", "Cartographie, fiches processus, grilles fournisseurs, trames 8D", "Écouter le client, arbitrer un écart de conformité produit"),
        ("C. Environnement — ISO 14001", "5", "Registre aspects/impacts, veille, plan d'urgence, filières déchets", "Engager une dépense environnementale"),
        ("D. Santé-sécurité — ISO 45001", "5", "Trames d'évaluation des risques, plan de prévention, grilles d'enquête", "Évaluer le risque à la place de l'employeur, signer le DUERP"),
        ("E. Actifs — ISO 55001", "5", "Registre d'actifs, grilles de criticité, trames LCC, plan de maintenance", "Arbitrer un investissement patrimonial"),
        ("F. Énergie — ISO 50001", "5", "Revue énergétique, fiches EnPI, plan de M&V, plan d'action", "Valider une économie réelle sur site"),
        ("G. Continuité — ISO 22301", "5", "BIA, trames de BCP, scénarios d'exercice, cartographie des dépendances", "Déclencher une cellule de crise"),
        ("H. Anti-corruption — ISO 37001", "4", "Cartographie du risque, questionnaires tiers, registre cadeaux, procédure d'alerte", "Instruire une alerte, décider d'une sanction"),
        ("I. Conformité — ISO 37301", "4", "Registre des obligations, matrice de risque de non-conformité, plan de formation", "Interpréter une obligation juridique"),
        ("J. SMS aérodrome", "17", "Manuel SMS, registre des dangers, trames d'évaluation, tableau SPI, trames d'enquête, dossiers MOC", "Accepter un risque, fermer une piste, signer une évaluation"),
        ("K. Carbone et ACA", "12", "Méthode de calcul, inventaire, dossier ACA, trajectoire chiffrée, fiches leviers", "Certifier un chiffre, acheter des crédits, engager un investissement"),
        ("L. Savoir-être", "7", "Grilles d'observation, trames d'entretien, supports de formation", "Porter la posture et l'autorité au nom du titulaire"),
        ("M. Clause 7.2 transverse", "14", "Analyse des besoins, référentiel de prérequis, dossier de preuves, plan d'efficacité", "Nommer, évaluer et sanctionner les personnes"),
        ("TOTAL", "99", "—", "—"),
    ],
    [4.0, 1.2, 7.3, 8.6], font_size=8.5
)

# ============================ 3 CASSETTES ============================
add_para("Trois casquettes opérationnelles", size=12, bold=True, color="1F3864", after=4)
add_table(
    ["Casquette", "Compétences dominantes", "Livrables types"],
    [
        ("Architecte du système intégré", "A01-A12, B, C, D, E, F, G, H, I, M", "Architecture documentaire, table HLS, analyse d'écarts, programme d'audit, revue de direction, matrice 7.2"),
        ("Ingénieur SMS aérodrome", "J01-J17 + D03, D04, G03, G04", "Manuel SMS, registre des dangers, évaluations, SPI, trames d'enquête, dossiers MOC, plan d'exercices"),
        ("Ingénieur carbone et ACA", "K01-K12 + F01-F05, C01-C02", "Note de périmètre, inventaire, dossier ACA, trajectoire, MACC, politique de compensation, dossier de vérification"),
    ],
    [4.0, 6.5, 11.0], font_size=8.5
)

# ============================ PERMISSIONS ============================
add_para("Permissions — ce que l'agent peut faire", size=12, bold=True, color="1F3864", after=4)
add_para("Peut : lire les documents internes et les données agrégées ; conduire des entretiens ; rédiger en brouillon ; calculer et consolider des données tracées ; alerter sur un écart ; préparer les dossiers d'audit, d'inspection et de vérification.",
         size=8.5, after=2)
add_para("Ne peut pas : publier, signer ou approuver ; supprimer un document validé ; attribuer un niveau de compétence sans preuve ; accepter un risque de sécurité ni valider un MOC ; certifier un chiffre d'émissions ni acheter des crédits ; contacter autorité, certificateur ou tiers sans mandat.",
         size=8.5, after=2)

doc.save(OUT)
print("Saved:", OUT)
